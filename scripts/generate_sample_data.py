"""Generate the sample documents used for demos, tests and the evaluation set.

All content is fictional ("Nusantara Digital", a made-up company) so it is
safe to publish. Run from the repository root:

    cd backend && uv run python ../scripts/generate_sample_data.py

Outputs (in sample_data/):
    hr/employee_handbook.pdf            native-text PDF with headings and paragraphs
    legal/service_agreement_template.docx  DOCX with title, headings, paragraphs, table
    legal/scanned_nda.pdf               image-only PDF (simulated scan) -> needs OCR
    invalid/corrupted.pdf               broken PDF  -> "could not be read" error
    invalid/corrupted.docx              broken DOCX -> "could not be read" error
    invalid/blank.pdf                   valid PDF with no text at all -> "no text" error
    invalid/empty.pdf                   0-byte file -> rejected at upload
"""

from __future__ import annotations

import random
from pathlib import Path

import docx
import reportlab
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parent.parent / "sample_data"
FONT_DIR = Path(reportlab.__file__).parent / "fonts"

# --------------------------------------------------------------------------
# 1. Native-text PDF: Employee Handbook (HR agent)
# --------------------------------------------------------------------------

HANDBOOK = [
    ("h1", "Nusantara Digital - Employee Handbook"),
    ("p", "Version 3.2, effective 1 January 2026. Owner: People & Culture team."),
    ("h2", "1. Welcome and Purpose"),
    ("p", "This handbook summarises the policies that apply to all permanent employees of "
          "Nusantara Digital. Contractors and interns follow the separate Contractor Guide. "
          "Where this handbook and an individual employment contract differ, the contract prevails."),
    ("h2", "2. Working Hours"),
    ("p", "Standard working hours are 40 hours per week, Monday to Friday. Core collaboration "
          "hours are 10:00 to 15:00 Western Indonesia Time (WIB); outside core hours employees "
          "may arrange their schedule flexibly with their manager."),
    ("p", "Overtime must be approved in advance by the direct manager. Approved overtime is "
          "compensated as time off in lieu at a rate of 1.5 hours for every hour worked."),
    ("h2", "3. Remote and Hybrid Work"),
    ("p", "Employees may work remotely up to 3 days per week. Fully remote arrangements require "
          "approval from the department head and People & Culture. The company provides a one-time "
          "home office allowance of IDR 2,500,000 for ergonomic equipment, claimable within the "
          "first 90 days of employment."),
    ("break", ""),
    ("h2", "4. Leave Policy"),
    ("h3", "4.1 Annual Leave"),
    ("p", "Permanent employees are entitled to 18 working days of paid annual leave per calendar "
          "year, accrued monthly at 1.5 days per month. Up to 5 unused days may be carried over to "
          "the next year and must be used before 31 March; remaining days expire."),
    ("p", "Annual leave requests must be submitted through the HR portal at least 7 days in "
          "advance for absences longer than 2 days."),
    ("h3", "4.2 Sick Leave"),
    ("p", "Employees receive up to 14 days of paid sick leave per year. A doctor's certificate is "
          "required for sick leave longer than 2 consecutive days."),
    ("h3", "4.3 Parental Leave"),
    ("p", "Maternity leave is 3 months (approximately 13 weeks) at full pay. Paternity leave is "
          "10 working days at full pay, to be taken within 60 days of the child's birth or adoption."),
    ("h3", "4.4 Other Leave"),
    ("p", "Employees may take 3 days of paid bereavement leave for the death of an immediate family "
          "member, and 2 days of paid leave for their own wedding."),
    ("break", ""),
    ("h2", "5. Expenses and Reimbursement"),
    ("p", "Business expenses must be submitted with receipts within 30 days of being incurred. "
          "Claims are reimbursed in the next monthly payroll. Meals during business travel are "
          "reimbursed up to IDR 300,000 per day. Economy class is required for flights under 6 hours."),
    ("h2", "6. Learning and Development"),
    ("p", "Each employee has an annual learning budget of IDR 7,500,000 for courses, certifications "
          "and conferences. Requests above IDR 2,000,000 require manager approval."),
    ("h2", "7. Performance Reviews"),
    ("p", "Performance reviews take place twice a year, in June and December. Salary adjustments "
          "are decided in the December cycle and take effect on 1 January."),
    ("h2", "8. Code of Conduct"),
    ("p", "Employees must avoid conflicts of interest and disclose any outside employment to "
          "People & Culture. Gifts from vendors worth more than IDR 500,000 must be declined or "
          "reported. Harassment of any kind is not tolerated and can be reported confidentially to "
          "ethics@nusantara-digital.example."),
]


def build_handbook_pdf(path: Path) -> None:
    styles = getSampleStyleSheet()
    style_map = {
        "h1": styles["Title"],
        "h2": styles["Heading2"],
        "h3": styles["Heading3"],
        "p": ParagraphStyle("body", parent=styles["BodyText"], fontSize=10.5, leading=15),
    }
    story = []
    for kind, text in HANDBOOK:
        if kind == "break":
            story.append(PageBreak())
            continue
        story.append(Paragraph(text, style_map[kind]))
        story.append(Spacer(1, 0.2 * cm))

    def footer(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(2 * cm, 1.2 * cm, "Nusantara Digital - Employee Handbook v3.2")
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {doc.page}")

    doc = SimpleDocTemplate(
        str(path), pagesize=A4, title="Employee Handbook", author="Nusantara Digital",
        leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
    )
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


# --------------------------------------------------------------------------
# 2. DOCX: Service Agreement template (Legal agent)
# --------------------------------------------------------------------------


def build_service_agreement_docx(path: Path) -> None:
    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)

    title = d.add_heading("Master Service Agreement - Template", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph(
        "This Master Service Agreement (the \"Agreement\") is entered into between Nusantara "
        "Digital (the \"Provider\") and the client named in the Order Form (the \"Client\"). "
        "Template reference: LEGAL-MSA-2026-01."
    )

    d.add_heading("1. Definitions", level=1)
    d.add_paragraph(
        "\"Services\" means the software development and support services described in each "
        "Statement of Work (SOW). \"Deliverables\" means all work product delivered to the Client "
        "under a SOW. \"Confidential Information\" has the meaning given in Section 6."
    )

    d.add_heading("2. Term", level=1)
    d.add_paragraph(
        "The Agreement starts on the Effective Date and continues for an initial term of 24 months. "
        "It renews automatically for successive 12-month periods unless either party gives written "
        "notice of non-renewal at least 60 days before the end of the current term."
    )

    d.add_heading("3. Fees and Payment", level=1)
    d.add_paragraph(
        "The Client pays the fees set out in the fee schedule below. Invoices are issued monthly in "
        "arrears and are payable within 30 days of the invoice date."
    )
    table = d.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    header[0].text, header[1].text, header[2].text = "Service tier", "Monthly fee", "Response time (SLA)"
    for tier, fee, sla in [
        ("Basic", "IDR 25,000,000", "Next business day"),
        ("Professional", "IDR 60,000,000", "8 business hours"),
        ("Enterprise", "IDR 150,000,000", "2 hours, 24/7"),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = tier, fee, sla
    d.add_paragraph(
        "Late payments accrue interest at 1% per month on the overdue amount. The Provider may "
        "suspend the Services if an invoice remains unpaid for more than 45 days after written notice."
    )

    d.add_heading("4. Termination", level=1)
    d.add_heading("4.1 Termination for Convenience", level=2)
    d.add_paragraph(
        "Either party may terminate this Agreement for convenience by giving 90 days' prior written "
        "notice. Fees for Services performed up to the termination date remain payable."
    )
    d.add_heading("4.2 Termination for Cause", level=2)
    d.add_paragraph(
        "Either party may terminate immediately by written notice if the other party commits a "
        "material breach and fails to remedy it within 30 days of receiving notice of the breach, "
        "or becomes insolvent."
    )

    d.add_heading("5. Intellectual Property", level=1)
    d.add_paragraph(
        "Upon full payment, ownership of all Deliverables transfers to the Client. The Provider "
        "retains ownership of its pre-existing tools and libraries and grants the Client a "
        "perpetual, royalty-free licence to use them as part of the Deliverables."
    )

    d.add_heading("6. Confidentiality", level=1)
    d.add_paragraph(
        "Each party shall protect the other party's Confidential Information using at least "
        "reasonable care and shall not disclose it to third parties. These obligations survive for "
        "5 years after termination of the Agreement."
    )

    d.add_heading("7. Limitation of Liability", level=1)
    d.add_paragraph(
        "Each party's total liability under this Agreement is limited to the fees paid by the Client "
        "in the 12 months preceding the claim. Neither party is liable for indirect or "
        "consequential damages, including lost profits."
    )

    d.add_heading("8. Governing Law and Disputes", level=1)
    d.add_paragraph(
        "This Agreement is governed by the laws of the Republic of Indonesia. Disputes that cannot "
        "be resolved amicably within 30 days shall be settled by arbitration at the Indonesian "
        "National Board of Arbitration (BANI) in Jakarta."
    )
    d.save(str(path))


# --------------------------------------------------------------------------
# 3. Image-only PDF: scanned NDA (Legal agent) -- requires OCR
# --------------------------------------------------------------------------

NDA_LINES = [
    ("title", "MUTUAL NON-DISCLOSURE AGREEMENT"),
    ("small", "Nusantara Digital  -  Form NDA-07  -  Signed copy (scan)"),
    ("gap", ""),
    ("head", "1. Purpose"),
    ("body", "The parties wish to exchange confidential information to evaluate"),
    ("body", "a potential business partnership (the \"Purpose\")."),
    ("gap", ""),
    ("head", "2. Duration"),
    ("body", "This NDA remains in force for 3 years from the date of signature."),
    ("body", "Confidentiality obligations continue for 2 years after it expires."),
    ("gap", ""),
    ("head", "3. Exclusions"),
    ("body", "Information is not confidential if it is publicly available, was"),
    ("body", "already known to the recipient, or is independently developed."),
    ("gap", ""),
    ("head", "4. Return of Materials"),
    ("body", "Within 14 days of a written request, the recipient must return or"),
    ("body", "destroy all confidential materials and confirm this in writing."),
    ("gap", ""),
    ("head", "5. Remedies"),
    ("body", "A breach entitles the disclosing party to seek injunctive relief and"),
    ("body", "liquidated damages of IDR 500,000,000 per breach."),
    ("gap", ""),
    ("head", "6. Governing Law"),
    ("body", "This NDA is governed by the laws of the Republic of Indonesia."),
    ("gap", ""),
    ("gap", ""),
    ("body", "Signed in Jakarta on 12 March 2026."),
]


def build_scanned_nda_pdf(path: Path) -> None:
    rng = random.Random(42)
    width, height = 1654, 2339  # A4 at 200 DPI
    page = Image.new("L", (width, height), 246)
    draw = ImageDraw.Draw(page)
    fonts = {
        "title": ImageFont.truetype(str(FONT_DIR / "VeraBd.ttf"), 46),
        "head": ImageFont.truetype(str(FONT_DIR / "VeraBd.ttf"), 34),
        "body": ImageFont.truetype(str(FONT_DIR / "Vera.ttf"), 32),
        "small": ImageFont.truetype(str(FONT_DIR / "VeraIt.ttf"), 26),
    }
    y = 170
    for kind, text in NDA_LINES:
        if kind == "gap":
            y += 34
            continue
        font = fonts[kind]
        x = 150 if kind != "title" else (width - draw.textlength(text, font=font)) / 2
        draw.text((x, y), text, fill=rng.randint(20, 45), font=font)
        y += int(font.size * 1.55)

    # Signature scribble + stamp so it looks like a signed scan.
    for _ in range(3):
        pts = [(1050 + i * 18, 2050 + rng.randint(-22, 22)) for i in range(18)]
        draw.line(pts, fill=40, width=3)
    draw.ellipse((300, 1960, 520, 2180), outline=90, width=6)
    draw.text((338, 2052), "APPROVED", fill=90, font=fonts["small"])

    # Scanner artefacts: slight skew, blur, speckle noise.
    page = page.rotate(0.7, resample=Image.BICUBIC, fillcolor=246, expand=False)
    page = page.filter(ImageFilter.GaussianBlur(0.6))
    pixels = page.load()
    for _ in range(9000):
        px, py = rng.randrange(width), rng.randrange(height)
        pixels[px, py] = rng.randint(120, 200)
    page.convert("RGB").save(path, "PDF", resolution=200.0)


# --------------------------------------------------------------------------
# 4. Invalid files for error-handling demos
# --------------------------------------------------------------------------


def build_invalid_files(folder: Path) -> None:
    (folder / "corrupted.pdf").write_bytes(
        b"%PDF-1.7\n1 0 obj << /Type /Catalog /Pages 2 0 R >>\n" + bytes(range(256)) * 8
    )
    (folder / "corrupted.docx").write_bytes(b"PK\x03\x04" + b"this is not a real zip archive" * 20)
    (folder / "empty.pdf").write_bytes(b"")
    blank = Image.new("RGB", (1240, 1754), "white")
    blank.save(folder / "blank.pdf", "PDF", resolution=150.0)


def main() -> None:
    for sub in ("hr", "legal", "invalid"):
        (ROOT / sub).mkdir(parents=True, exist_ok=True)
    build_handbook_pdf(ROOT / "hr" / "employee_handbook.pdf")
    build_service_agreement_docx(ROOT / "legal" / "service_agreement_template.docx")
    build_scanned_nda_pdf(ROOT / "legal" / "scanned_nda.pdf")
    build_invalid_files(ROOT / "invalid")
    for f in sorted(ROOT.rglob("*.*")):
        print(f"{f.relative_to(ROOT.parent)}  ({f.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
